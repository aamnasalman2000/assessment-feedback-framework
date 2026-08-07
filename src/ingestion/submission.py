from __future__ import annotations

import mimetypes
from pathlib import Path
from typing import Any


IGNORED_NAMES = {
    ".DS_Store",
    "Thumbs.db",
}


TEXT_EXTENSIONS = {
    ".lean",
    ".pl",
    ".txt",
    ".md",
    ".owl",
    ".owx",
    ".rdf",
    ".xml",
}


ONTOLOGY_EXTENSIONS = {
    ".owl",
    ".owx",
    ".rdf",
    ".xml",
}


def discover_submission_files(
    student_directory: Path,
) -> list[Path]:
    if not student_directory.exists():
        raise FileNotFoundError(
            "Student directory not found: "
            f"{student_directory}"
        )

    files = [
        path
        for path in student_directory.rglob("*")
        if (
            path.is_file()
            and path.name not in IGNORED_NAMES
        )
    ]

    return sorted(
        files,
        key=lambda path: str(path).lower(),
    )


def _is_prolog_filename(
    path: Path,
) -> bool:
    """
    Identify files that contain Prolog source even when their extension
    is .txt or .docx.
    """
    filename = path.name.lower()

    return any(
        marker in filename
        for marker in (
            "_pl",
            "-pl",
            "prolog",
        )
    )


def infer_artifact_type(
    path: Path,
    *,
    assessment_id: str,
) -> str:
    """
    Infer the assessment role of a submitted file.

    Assessment 2 contains ontology, report, and Prolog artefacts. Some
    Prolog submissions use .txt or .docx rather than .pl.
    """
    suffix = path.suffix.lower()

    if suffix == ".lean":
        return "lean_source"

    if suffix in ONTOLOGY_EXTENSIONS:
        return "owl_ontology"

    if suffix == ".pl":
        return "prolog_source"

    if assessment_id == "assessment_2":
        if suffix == ".txt":
            return "prolog_source"

        if (
            suffix == ".docx"
            and _is_prolog_filename(path)
        ):
            return "prolog_source"

        if suffix == ".pdf":
            return "ontology_report"

    if suffix == ".pdf":
        return "document"

    if suffix == ".docx":
        return "document"

    if suffix in {
        ".txt",
        ".md",
    }:
        return "text_document"

    return "unknown_artifact"


def read_text_content(
    path: Path,
) -> tuple[str, str]:
    """
    Read a text-based source file with conservative encoding fallbacks.
    """
    encodings = (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    )

    for encoding in encodings:
        try:
            return (
                path.read_text(encoding=encoding),
                encoding,
            )
        except UnicodeDecodeError:
            continue

    return (
        path.read_text(
            encoding="utf-8",
            errors="replace",
        ),
        "utf-8",
    )


def extract_pdf_text(
    path: Path,
) -> tuple[str, str]:
    """
    Extract text from a PDF report.

    Requires:
        pip install pypdf
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF processing requires the 'pypdf' "
            "package. Install it with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))

    page_texts: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = page.extract_text() or ""

        page_texts.append(
            f"[PAGE {page_number}]\n{text.strip()}"
        )

    return (
        "\n\n".join(page_texts).strip(),
        "utf-8",
    )


def extract_docx_text(
    path: Path,
) -> tuple[str, str]:
    """
    Extract text from a DOCX document.

    Requires:
        pip install python-docx
    """
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX processing requires the "
            "'python-docx' package. Install it with: "
            "pip install python-docx"
        ) from exc

    document = Document(str(path))

    content: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.rstrip()

        if text:
            content.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                content.append(
                    "\t".join(cells)
                )

    return (
        "\n".join(content).strip(),
        "utf-8",
    )


def read_raw_content(
    path: Path,
) -> tuple[str, str]:
    """
    Extract textual content from supported source and document formats.
    """
    suffix = path.suffix.lower()

    if suffix in TEXT_EXTENSIONS:
        return read_text_content(path)

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    # The schema requires raw_content to be a string. Unsupported binary
    # artefacts retain an empty string and are handled later as unsupported.
    return "", "utf-8"


def infer_associations(
    *,
    artifact_type: str,
    assessment_id: str,
) -> dict[str, Any] | None:
    """
    Add high-level routing information when it can be determined from the
    artefact type without guessing task-level associations.
    """
    if assessment_id != "assessment_2":
        return None

    if artifact_type in {
        "owl_ontology",
        "ontology_report",
    }:
        return {
            "part_ids": [
                "part_1_ontology",
            ],
            "source": "inferred",
        }

    if artifact_type == "prolog_source":
        return {
            "part_ids": [
                "part_2_prolog_modal_logic",
            ],
            "source": "inferred",
        }

    return None


def build_file_record(
    path: Path,
    student_directory: Path,
    index: int,
    *,
    assessment_id: str,
) -> dict[str, Any]:
    media_type, _ = mimetypes.guess_type(
        path.name
    )

    artifact_type = infer_artifact_type(
        path,
        assessment_id=assessment_id,
    )

    raw_content, content_encoding = (
        read_raw_content(path)
    )

    suffix = path.suffix.lower().lstrip(".")
    file_format = suffix or "unknown"

    record: dict[str, Any] = {
        "artifact_id": f"artifact_{index:03d}",
        "artifact_type": artifact_type,
        "source": {
            "filename": path.name,
            "format": file_format,
            "media_type": media_type,
            "path": (
                path.relative_to(
                    student_directory
                ).as_posix()
            ),
        },
        "content": {
            "raw_content": raw_content,
            "content_encoding": (
                content_encoding
            ),
        },
    }

    associations = infer_associations(
        artifact_type=artifact_type,
        assessment_id=assessment_id,
    )

    if associations is not None:
        record["associations"] = associations

    return record


def build_submission(
    *,
    assessment_id: str,
    student_id: str,
    student_directory: Path,
    schema_version: str = "1.0",
) -> dict[str, Any]:
    files = discover_submission_files(
        student_directory
    )

    if not files:
        raise ValueError(
            "No submission files found in "
            f"{student_directory}"
        )

    artifacts = [
        build_file_record(
            path=path,
            student_directory=student_directory,
            index=index,
            assessment_id=assessment_id,
        )
        for index, path in enumerate(
            files,
            start=1,
        )
    ]

    return {
        "schema_version": schema_version,
        "submission_id": (
            f"{assessment_id}_{student_id}"
        ),
        "student_id": student_id,
        "assessment_id": assessment_id,
        "artifacts": artifacts,
    }